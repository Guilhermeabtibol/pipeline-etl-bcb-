import os
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_to_excel(
    df: pd.DataFrame, output_path: str = "cotacoes_dolar.xlsx"
) -> str:
    """Exporta o DataFrame tratado para Excel com formatação visual profissional."""
    if df.empty:
        raise ValueError(
            "DataFrame vazio. Nenhum dado disponível para exportar."
        )

    # Exportação inicial via Pandas com engine openpyxl
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Cotações Dólar", index=False)

        # Acesso ao workbook e worksheet para estilização
        workbook = writer.book
        worksheet = writer.sheets["Cotações Dólar"]

        # Estilos: Cabeçalho com fundo escuro e texto em negrito/branco
        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        align_center = Alignment(horizontal="center", vertical="center")

        # Aplica estilo no cabeçalho
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center

        # Ajuste automático de largura de colunas e alinhamentos
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)

            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))

                # Ajuste de alinhamento por linha de dados
                if cell.row > 1:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = (
                            "R$ #,##0.00"  # Formato Moeda BRL
                        )
                    elif isinstance(cell.value, str) and len(val_str) == 10:
                        cell.alignment = align_center  # Formato Data YYYY-MM-DD

            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    return os.path.abspath(output_path)